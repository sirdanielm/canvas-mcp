"""Explicit, local, copy-on-write PRIVATE pseudonymous archival exports.

These values-only copies are never operational workbook/baseline inputs.
"""

from __future__ import annotations

import csv
import hashlib
import html
import importlib.util
import io
import json
import os
import re
import secrets
import sys
import unicodedata
import xml.etree.ElementTree as ET
from collections import defaultdict
from collections.abc import Callable
from pathlib import Path
from types import ModuleType
from typing import Any
from urllib.parse import unquote
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

PIN = re.compile(r"[0-9]{4}\Z")
EMAIL = re.compile(
    r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+"
)
NAME_HEADERS = {
    "Student Name",
    "Student",
    "Name",
    "student_name",
    "Student name",
    "Canvas Name",
    "IC Name",
    "Roster Name",
}
EMAIL_HEADERS = {"Email", "Student Email", "Email Address", "student_email", "email"}
PIN_HEADERS = {"Student Number", "Student PIN", "PIN"}
RESERVED = {"Points possible", "Assignment IDs"}


class ExportHeld(ValueError):
    """Privacy-safe error code only; never include source values or paths."""


class ExportPublicationUncertain(RuntimeError):
    """Publication began; an archive may exist and must not be blindly retried."""


def require(ok: object, code: str) -> None:
    if not ok:
        raise ExportHeld(code)


def fingerprint(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def private_directory_identity(directory: Path) -> tuple[int, int]:
    require(directory.is_dir(), "OUTPUT_DIRECTORY_REQUIRED")
    require(
        all(not p.is_symlink() for p in (directory, *directory.parents)),
        "OUTPUT_SYMLINK_NOT_ALLOWED",
    )
    require(
        not any((p / ".git").exists() for p in (directory, *directory.parents)),
        "ARCHIVE_INSIDE_GIT_NOT_ALLOWED",
    )
    info = directory.stat()
    require(info.st_mode & 0o077 == 0, "PRIVATE_OUTPUT_DIRECTORY_REQUIRED")
    return info.st_dev, info.st_ino


def open_xlsx(path: Path, *, raw: bytes | None = None) -> Workbook:
    # Every parser pass must use the same bytes as the source hash receipt.
    raw = path.read_bytes() if raw is None else raw
    with ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        require(
            len({entry.filename for entry in entries}) == len(entries),
            "DUPLICATE_WORKBOOK_MEMBER",
        )
        require(
            len(entries) <= 1000 and sum(x.file_size for x in entries) <= 64_000_000,
            "WORKBOOK_TOO_LARGE",
        )
        require(
            all(x.filename.endswith((".xml", ".rels", ".vml")) for x in entries),
            "UNSUPPORTED_EMBEDDED_CONTENT",
        )
        require(
            not any(
                any(
                    part in x.filename
                    for part in (
                        "charts/",
                        "externalLinks/",
                        "pivot",
                        "embeddings/",
                        "vbaProject",
                        "activeX/",
                        "ctrlProps/",
                        "customXml/",
                    )
                )
                for x in entries
            ),
            "UNSUPPORTED_WORKBOOK_FEATURE",
        )
        for info in entries:
            if info.filename.endswith(".vml"):
                root = ET.fromstring(archive.read(info))
                shapes = root.findall(".//{urn:schemas-microsoft-com:vml}shape")
                require(bool(shapes), "UNSUPPORTED_DRAWING")
                for shape in shapes:
                    data = shape.find(
                        "{urn:schemas-microsoft-com:office:excel}ClientData"
                    )
                    require(
                        data is not None and data.get("ObjectType") == "Note",
                        "UNSUPPORTED_DRAWING",
                    )
            elif "drawings/" in info.filename:
                require(
                    len(ET.fromstring(archive.read(info))) == 0, "UNSUPPORTED_DRAWING"
                )
    book = load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    require(sum(len(s._cells) for s in book) <= 200_000, "WORKBOOK_TOO_LARGE")
    return book


def roster_mapping(
    path: Path, sheet_name: str | None = None, *, raw: bytes | None = None
) -> tuple[dict[str, set[str]], set[str]]:
    book = open_xlsx(path, raw=raw)
    names = (
        [sheet_name]
        if sheet_name
        else [n for n in ("Student Info", "Student Numbers") if n in book.sheetnames]
    )
    require(
        len(names) == 1 and names[0] in book.sheetnames, "EXACT_ROSTER_SHEET_REQUIRED"
    )
    sheet = book[names[0]]
    rows = []
    for row_number in range(1, min(sheet.max_row, 10) + 1):
        values = [
            sheet.cell(row_number, c).value for c in range(1, sheet.max_column + 1)
        ]
        hits = {
            key: [i for i, v in enumerate(values) if v in labels]
            for key, labels in (
                ("name", NAME_HEADERS),
                ("email", EMAIL_HEADERS),
                ("pin", PIN_HEADERS),
            )
        }
        if len(hits["name"]) >= 1 and len(hits["email"]) == 1 and len(hits["pin"]) == 1:
            rows.append((row_number, hits))
    require(len(rows) == 1, "ROSTER_HEADERS_NOT_UNIQUE")
    header, cols = rows[0]
    aliases: dict[str, set[str]] = defaultdict(set)
    pin_emails: dict[str, set[str]] = defaultdict(set)
    count = 0
    for row in sheet.iter_rows(min_row=header + 1):
        name_values = [row[c].value for c in cols["name"]]
        email, pin = (row[cols[k][0]].value for k in ("email", "pin"))
        if all(v in (None, "") for v in [*name_values, email, pin]):
            continue
        require(
            isinstance(pin, str) and PIN.fullmatch(pin), "FOUR_DIGIT_TEXT_PIN_REQUIRED"
        )
        roster_names = [n for n in name_values if n not in (None, "")]
        require(
            roster_names
            and all(isinstance(n, str) and n.strip() == n and n for n in roster_names),
            "ROSTER_NAME_INVALID",
        )
        require(
            isinstance(email, str) and EMAIL.fullmatch(email), "ROSTER_EMAIL_INVALID"
        )
        require(
            all(row[c].data_type != "f" for indices in cols.values() for c in indices),
            "ROSTER_FORMULA_NOT_ALLOWED",
        )
        assert isinstance(pin, str) and isinstance(email, str)
        for name in roster_names:
            assert isinstance(name, str)
            aliases[name].add(pin)
        aliases[email].add(pin)
        pin_emails[pin].add(email)
        count += 1
    require(
        count > 0 and all(len(v) == 1 for v in pin_emails.values()),
        "ROSTER_PIN_CONFLICT",
    )
    require(
        all(len(v) == 1 for k, v in aliases.items() if EMAIL.fullmatch(k)),
        "ROSTER_EMAIL_CONFLICT",
    )
    return dict(aliases), set(pin_emails)


def text_sanitizer(aliases: dict[str, set[str]]) -> Callable[[Any], Any]:
    # Exact complete identity strings only; no normalization, guesses or new PINs.
    expression = re.compile(
        r"(?<!\w)(?:"
        + "|".join(re.escape(k) for k in sorted(aliases, key=len, reverse=True))
        + r")(?!\w)"
    )

    residual_expression = re.compile(expression.pattern, re.IGNORECASE)

    def sanitize(value: Any) -> Any:
        if not isinstance(value, str):
            return value
        decoded = value
        for _ in range(4):
            next_decoded = html.unescape(unquote(decoded))
            if next_decoded == decoded:
                break
            require(
                not EMAIL.search(next_decoded)
                and not residual_expression.search(next_decoded),
                "ENCODED_IDENTITY_NOT_SUPPORTED",
            )
            decoded = next_decoded
        require(
            html.unescape(unquote(decoded)) == decoded,
            "ENCODING_DEPTH_NOT_SUPPORTED",
        )

        def replace(match: re.Match[str]) -> str:
            choices = aliases[match.group()]
            require(len(choices) == 1, "AMBIGUOUS_IDENTITY_REFERENCE")
            return next(iter(choices))

        def unique_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
            result: dict[str, Any] = {}
            for key, item in pairs:
                require(key not in result, "DUPLICATE_JSON_KEY")
                result[key] = item
            return result

        try:
            nested = json.loads(value, object_pairs_hook=unique_object)
        except ExportHeld:
            raise
        except (ValueError, TypeError):
            nested = None
        if isinstance(nested, (dict, list)):

            def walk(item: Any) -> Any:
                if isinstance(item, dict):
                    keys = [sanitize(k) for k in item]
                    require(len(set(keys)) == len(keys), "SANITIZED_JSON_KEY_CONFLICT")
                    return {
                        key: walk(v) for key, v in zip(keys, item.values(), strict=True)
                    }
                if isinstance(item, list):
                    return [walk(v) for v in item]
                return sanitize(item)

            value = json.dumps(walk(nested), ensure_ascii=False, separators=(",", ":"))
        elif isinstance(nested, str):
            # A JSON string is also an encoded identity carrier, even when it
            # is not wrapped in an object or array. Preserve its string type.
            value = json.dumps(sanitize(nested), ensure_ascii=False)

        # Check complete email tokens before name substitution. Word boundaries
        # alone can replace a known address inside a longer, unknown address.
        def replace_email(match: re.Match[str]) -> str:
            adjacent = (
                value[max(0, match.start() - 1) : match.start()]
                + value[match.end() : match.end() + 1]
            )
            require(
                not any(
                    c.isalnum() or c == "_" or unicodedata.category(c).startswith("M")
                    for c in adjacent
                ),
                "UNMAPPED_EMAIL_REMAINS",
            )
            choices = aliases.get(match.group(), set())
            require(bool(choices), "UNMAPPED_EMAIL_REMAINS")
            require(len(choices) == 1, "AMBIGUOUS_IDENTITY_REFERENCE")
            return next(iter(choices))

        without_emails = EMAIL.sub(replace_email, value)
        # Unsupported quoted, internationalized or local-only mailboxes must
        # not evade the ASCII mailbox matcher and remain in a successful copy.
        require("@" not in without_emails, "UNMAPPED_EMAIL_REMAINS")
        result = expression.sub(replace, without_emails)
        require(not EMAIL.search(result), "UNMAPPED_EMAIL_REMAINS")
        require(not residual_expression.search(result), "IDENTITY_REMAINS")
        return result

    return sanitize


def replace_columns(
    rows: list[list[Any]], aliases: dict[str, set[str]], pins: set[str]
) -> int:
    """Rows are mutable lists. Discover explicit labels in the first ten rows."""
    headers = []
    for number, row in enumerate(rows[:10]):
        columns = [
            i
            for i, value in enumerate(row)
            if isinstance(value, str) and value in NAME_HEADERS | EMAIL_HEADERS
        ]
        if columns:
            headers.append((number, columns))
    require(len(headers) <= 1, "IDENTITY_HEADERS_NOT_UNIQUE")
    if not headers:
        return 0
    header, columns = headers[0]
    replaced = 0
    for number, row in enumerate(rows[header + 1 :], header + 1):
        values = [row[c] if c < len(row) else None for c in columns]
        if all(v in (None, "") for v in values):
            continue
        # Canonical gradebook has three header rows: Student / Points / IDs.
        if number in (header + 1, header + 2) and all(
            v in (None, "") or v in RESERVED for v in values
        ):
            continue
        candidates = []
        for value in values:
            if value in (None, ""):
                continue
            require(isinstance(value, str), "IDENTITY_NOT_TEXT")
            assert isinstance(value, str)
            choices = {value} if value in pins else aliases.get(value, set())
            # Canonical GET/Edit labels are exact name + newline + sections.
            if not choices and header == 2 and columns == [1] and "\n" in value:
                choices = aliases.get(value.split("\n", 1)[0], set())
            require(bool(choices), "UNMAPPED_IDENTITY")
            candidates.append(choices)
        matching = set.intersection(*candidates)
        require(len(matching) == 1, "IDENTITY_CONFLICT_OR_AMBIGUITY")
        pin = next(iter(matching))
        for c in columns:
            while len(row) <= c:
                row.append(None)
            row[c] = pin
        replaced += 1
    for c in columns:
        rows[header][c] = (
            "Student PIN" if rows[header][c] in NAME_HEADERS else "Student PIN (email)"
        )
    return replaced


def sanitize_export(
    source: Path,
    central_roster: Path,
    output: Path,
    roster_sheet: str | None = None,
    *,
    verified_mapping: tuple[dict[str, set[str]], set[str]] | None = None,
    expected_roster_hash: str | None = None,
    authority_check: Callable[[], None] | None = None,
) -> dict[str, Any]:
    source, central_roster, output = map(Path, (source, central_roster, output))
    require(
        all(p.is_absolute() for p in (source, central_roster, output)),
        "ABSOLUTE_PATH_REQUIRED",
    )
    require(
        source.resolve() != output.resolve()
        and central_roster.resolve() != output.resolve(),
        "ORIGINAL_MUST_BE_PRESERVED",
    )
    require(not output.exists(), "OUTPUT_ALREADY_EXISTS")
    directory_identity = private_directory_identity(output.parent)
    require(
        source.suffix.lower() in (".xlsx", ".csv")
        and output.suffix.lower() == source.suffix.lower(),
        "FORMAT_NOT_SUPPORTED",
    )
    source_raw, roster_raw = source.read_bytes(), central_roster.read_bytes()
    before = hashlib.sha256(source_raw).hexdigest()
    roster_before = hashlib.sha256(roster_raw).hexdigest()
    require(
        expected_roster_hash is None or expected_roster_hash == roster_before,
        "SHARED_PIN_ROSTER_CHANGED",
    )
    aliases, pins = (
        verified_mapping
        if verified_mapping is not None
        else roster_mapping(central_roster, roster_sheet, raw=roster_raw)
    )
    sanitize = text_sanitizer(aliases)
    count = 0
    if source.suffix.lower() == ".csv":
        try:
            csv_rows: list[list[Any]] = list(
                csv.reader(io.StringIO(source_raw.decode("utf-8-sig")), strict=True)
            )
        except csv.Error as error:
            raise ExportHeld("CSV_INVALID") from error
        require(len(csv_rows) <= 200_000, "EXPORT_TOO_LARGE")
        count += replace_columns(csv_rows, aliases, pins)
        csv_values = [[sanitize(v) for v in row] for row in csv_rows]
        # CSV cannot declare literal string cells; prevent spreadsheet formulas.
        require(
            not any(
                v.startswith(("=", "+", "-", "@", "\t", "\r"))
                for r in csv_values
                for v in r
            ),
            "CSV_FORMULA_NOT_SUPPORTED",
        )
        csv_stream = io.StringIO(newline="")
        csv.writer(csv_stream).writerows(csv_values)
        payload = csv_stream.getvalue().encode("utf-8")
    else:
        book = open_xlsx(source, raw=source_raw)
        cached = load_workbook(io.BytesIO(source_raw), data_only=True, keep_links=False)
        clean = Workbook()
        active = clean.active
        assert active is not None
        clean.remove(active)
        for sheet in book:
            require(
                sheet.max_row * sheet.max_column <= 1_000_000,
                "SHEET_SHAPE_NOT_SUPPORTED",
            )
            rows: list[list[Any]] = []
            for source_row in sheet.iter_rows():
                values: list[Any] = []
                for cell in source_row:
                    value = (
                        cached[sheet.title][cell.coordinate].value
                        if cell.data_type == "f"
                        else cell.value
                    )
                    require(
                        cell.data_type != "f" or value is not None,
                        "FORMULA_CACHE_MISSING",
                    )
                    values.append(value)
                rows.append(values)
            count += replace_columns(rows, aliases, pins)
            title = sanitize(sheet.title)
            require(title not in clean.sheetnames, "SANITIZED_SHEET_TITLE_CONFLICT")
            dest = clean.create_sheet(title)
            dest.sheet_state = sheet.sheet_state
            for row_no, row in enumerate(rows, 1):
                for col_no, value in enumerate(row, 1):
                    if value is None:
                        continue
                    value = sanitize(value)
                    cell = dest.cell(row_no, col_no, value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                        if PIN.fullmatch(value):
                            cell.number_format = "@"
        stream = io.BytesIO()
        clean.save(stream)
        payload = stream.getvalue()
        # Rebuilt package drops all source comments, hyperlinks, names, metadata,
        # validation lists, drawings and cached shared strings, including hidden sheets.
        with ZipFile(io.BytesIO(payload)) as archive:
            for info in archive.infolist():
                text = archive.read(info).decode("utf-8")
                # XML namespaces contain no mailbox addresses. Unescape text first.
                require(
                    sanitize(html.unescape(text)) == html.unescape(text),
                    "OUTPUT_IDENTITY_REMAINS",
                )
    require(count > 0, "NO_IDENTITY_COLUMNS_FOUND")
    require(
        fingerprint(source) == before and fingerprint(central_roster) == roster_before,
        "INPUT_CHANGED",
    )
    if authority_check is not None:
        authority_check()
    directory_fd = os.open(output.parent, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    publication_attempted = False
    try:

        def check_directory() -> None:
            info = os.fstat(directory_fd)
            require(
                (info.st_dev, info.st_ino) == directory_identity
                and private_directory_identity(output.parent) == directory_identity,
                "OUTPUT_DIRECTORY_CHANGED",
            )
            require(info.st_mode & 0o077 == 0, "PRIVATE_OUTPUT_DIRECTORY_REQUIRED")

        check_directory()
        temporary = ".pin-export-" + secrets.token_hex(16)
        fd = os.open(
            temporary,
            os.O_CREAT | os.O_EXCL | os.O_WRONLY,
            0o600,
            dir_fd=directory_fd,
        )
        temporary_info = os.fstat(fd)

        def owns_temporary() -> bool:
            try:
                info = os.stat(temporary, dir_fd=directory_fd, follow_symlinks=False)
            except FileNotFoundError:
                return False
            return (info.st_dev, info.st_ino) == (
                temporary_info.st_dev,
                temporary_info.st_ino,
            )

        try:
            with os.fdopen(fd, "wb") as handle:
                handle.write(payload)
                handle.flush()
                os.fsync(handle.fileno())
            if authority_check is not None:
                authority_check()
            require(
                fingerprint(source) == before
                and fingerprint(central_roster) == roster_before,
                "INPUT_CHANGED",
            )
            check_directory()
            require(owns_temporary(), "TEMPORARY_OUTPUT_CHANGED")
            # Descriptor-relative creation cannot be redirected by a directory
            # replacement. Hard linking remains exclusive: no archive overwrite.
            publication_attempted = True
            try:
                os.link(
                    temporary,
                    output.name,
                    src_dir_fd=directory_fd,
                    dst_dir_fd=directory_fd,
                    follow_symlinks=False,
                )
            except FileExistsError as error:
                publication_attempted = False
                raise ExportHeld("OUTPUT_ALREADY_EXISTS") from error
            os.fsync(directory_fd)
            check_directory()
        finally:
            if owns_temporary():
                os.unlink(temporary, dir_fd=directory_fd)
            elif publication_attempted:
                raise ExportPublicationUncertain("ARCHIVE_PUBLICATION_UNCERTAIN")
            else:
                raise ExportHeld("TEMPORARY_OUTPUT_CHANGED")
    except Exception as error:
        if publication_attempted:
            raise ExportPublicationUncertain("ARCHIVE_PUBLICATION_UNCERTAIN") from error
        raise
    finally:
        try:
            os.close(directory_fd)
        except OSError as error:
            if publication_attempted:
                raise ExportPublicationUncertain(
                    "ARCHIVE_PUBLICATION_UNCERTAIN"
                ) from error
            raise
    return {
        "status": "PRIVATE_PSEUDONYMOUS_ARCHIVE",
        "identity_rows_replaced": count,
        "source_sha256": before,
        "roster_sha256": roster_before,
        "output_sha256": hashlib.sha256(payload).hexdigest(),
        "originals_preserved": True,
        "live_calls": 0,
    }


def default_shared_pin_tool(repository: Path) -> Path:
    """Resolve the canonical sibling even when this renderer runs in a worktree."""
    marker = repository / ".git"
    canonical = repository
    if marker.is_file():
        text = marker.read_text(encoding="utf-8").strip()
        require(text.startswith("gitdir: "), "GIT_WORKTREE_METADATA_INVALID")
        gitdir = (repository / text[8:]).resolve()
        common = gitdir / "commondir"
        require(common.is_file(), "GIT_WORKTREE_METADATA_INVALID")
        canonical = (
            (gitdir / common.read_text(encoding="utf-8").strip()).resolve().parent
        )
    return canonical.parent / "LocalGrAss-github" / "scripts" / "export-student-pins.py"


def load_shared_pin_module(tool: Path) -> ModuleType:
    require(
        tool.is_absolute() and tool.is_file() and not tool.is_symlink(),
        "SHARED_PIN_TOOL_UNAVAILABLE",
    )
    package = tool.parent.parent / "localgrass"
    require(
        (package / "__init__.py").is_file() and (package / "student_pin.py").is_file(),
        "SHARED_PIN_TOOL_UNAVAILABLE",
    )
    namespace = (
        "_grass_shared_pins_" + hashlib.sha256(str(package).encode()).hexdigest()[:16]
    )
    spec = importlib.util.spec_from_file_location(
        namespace, package / "__init__.py", submodule_search_locations=[str(package)]
    )
    require(spec is not None and spec.loader is not None, "SHARED_PIN_TOOL_UNAVAILABLE")
    assert spec is not None and spec.loader is not None
    loaded_package = importlib.util.module_from_spec(spec)
    sys.modules[namespace] = loaded_package
    spec.loader.exec_module(loaded_package)
    spec = importlib.util.spec_from_file_location(
        namespace + ".student_pin", package / "student_pin.py"
    )
    require(spec is not None and spec.loader is not None, "SHARED_PIN_TOOL_UNAVAILABLE")
    assert spec is not None and spec.loader is not None
    shared = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = shared
    spec.loader.exec_module(shared)
    require(
        all(
            callable(getattr(shared, name, None))
            for name in ("binding", "read_table", "roster_index")
        ),
        "SHARED_PIN_CONTRACT_UNAVAILABLE",
    )
    return shared


def registry_bound_export(
    source: Path, output: Path, registry_path: Path, shared_tool: Path
) -> dict[str, Any]:
    """Use the shared authority/index; this module only renders a private archive."""
    shared = load_shared_pin_module(shared_tool)
    require(registry_path.is_absolute(), "ABSOLUTE_PATH_REQUIRED")
    registry_before = fingerprint(registry_path)
    try:
        registry, roster, raw, registry_hash = shared.binding(registry_path)
        require(
            registry.get("roster_sheet") == "Student Info",
            "CURRENT_ROSTER_SHEET_REQUIRED",
        )
        headers, rows = shared.read_table(roster, raw, registry["roster_sheet"])
        index, names = shared.roster_index(headers, rows, registry)
    except ExportHeld:
        raise
    except Exception as error:
        raise ExportHeld("SHARED_PIN_AUTHORITY_REJECTED") from error
    require(
        registry_hash == registry_before
        and fingerprint(registry_path) == registry_before,
        "SHARED_PIN_REGISTRY_CHANGED",
    )
    require(
        fingerprint(roster) == registry["roster_sha256"], "SHARED_PIN_ROSTER_CHANGED"
    )
    # Retain original spellings for exact substitution, obtaining every PIN
    # from the shared index. No second assignment/normalization policy here.
    aliases: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        if not any(value is not None and str(value).strip() for value in row.values()):
            continue
        email = row[registry.get("email_column", "Email")]
        require(isinstance(email, str), "SHARED_PIN_MAPPING_INVALID")
        pin = index[email.strip().casefold()]
        require(
            isinstance(pin, str) and PIN.fullmatch(pin), "SHARED_PIN_MAPPING_INVALID"
        )
        aliases[email].add(pin)
        for column in registry.get("name_columns", []):
            name = row[column]
            if name in (None, ""):
                continue
            require(
                isinstance(name, str)
                and pin in names.get(name.strip().casefold(), set()),
                "SHARED_PIN_MAPPING_INVALID",
            )
            # Exact spellings must retain the shared index's normalized
            # ambiguity; casing alone is not unique identity evidence.
            aliases[name].update(names[name.strip().casefold()])

    def authority_check() -> None:
        require(
            fingerprint(registry_path) == registry_hash, "SHARED_PIN_REGISTRY_CHANGED"
        )
        require(
            fingerprint(roster) == registry["roster_sha256"],
            "SHARED_PIN_ROSTER_CHANGED",
        )

    result = sanitize_export(
        source,
        roster,
        output,
        registry["roster_sheet"],
        verified_mapping=(dict(aliases), set(index.values())),
        expected_roster_hash=registry["roster_sha256"],
        authority_check=authority_check,
    )
    result.update(registry_sha256=registry_hash, shared_pin_authority=True)
    return result
